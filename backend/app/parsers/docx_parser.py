"""DOCX parser using python-docx.

Extracts text paragraph-by-paragraph. Each paragraph is assigned to a
logical "page" approximated by paragraph index (DOCX has no native pages).
"""

from __future__ import annotations

import logging
from pathlib import Path

from app.core.exceptions import MalformedFileError
from app.parsers.base import BaseParser, ParsedDocument, ParsedPage

logger = logging.getLogger(__name__)

# Approximate number of paragraphs treated as one "page" for metadata purposes
_PARAS_PER_PAGE = 30


class DocxParser(BaseParser):
    """Parse DOCX files using python-docx."""

    mime_types = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    )
    extensions = ("docx",)

    def parse(self, file_path: Path) -> ParsedDocument:
        """Extract paragraphs from the DOCX, grouped into logical pages.

        Raises:
            MalformedFileError: If the file cannot be opened.
        """
        try:
            from docx import Document as DocxDocument  # lazy import
        except ImportError as exc:
            raise RuntimeError(
                "python-docx is not installed. Run: pip install python-docx"
            ) from exc

        try:
            doc = DocxDocument(str(file_path))
        except Exception as exc:
            logger.debug("DOCX open failed for %s: %s", file_path.name, type(exc).__name__)
            raise MalformedFileError(
                message="The Word document could not be read. It may be corrupted."
            ) from exc

        # Collect all non-empty paragraph texts
        para_texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

        # Group into logical pages
        pages: list[ParsedPage] = []
        for page_idx, start in enumerate(range(0, max(len(para_texts), 1), _PARAS_PER_PAGE)):
            chunk = para_texts[start : start + _PARAS_PER_PAGE]
            pages.append(
                ParsedPage(
                    page_number=page_idx + 1,
                    text="\n".join(chunk),
                )
            )

        if not pages:
            pages = [ParsedPage(page_number=1, text="")]

        logger.debug(
            "DOCX parsed: %d paragraphs → %d logical pages, %d chars",
            len(para_texts),
            len(pages),
            sum(p.char_count for p in pages),
        )
        return ParsedDocument(
            pages=pages,
            total_pages=len(pages),
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
